#!/usr/bin/env python3
"""
maak_werkdocumenten.py
Genereert het materiaal voor les 06 "Mijn NovaDepot-map" (bestandsbeheer, NovaDepot, ORLO).

  WP1_Mijn-NovaDepot-map.docx   het werkdocument dat de leerling INLEVERT
  rommel/                       vier bestanden met slechte namen (opzettelijk!)
  NovaDepot-bestanden.zip       die vier bestanden, zonder mapniveau: dit hang je aan de opdracht

Gebruik:  python3 maak_werkdocumenten.py
Vereist:  python-docx, openpyxl, Pillow

LET OP bij aanpassen:
 - De slechte bestandsnamen zijn LESMATERIAAL, geen slordigheid: geen enkele naam zegt wat erin zit.
 - Elk bestand toont bovenaan in grote letters WAT het is en VAN WIE. Daaruit maakt de leerling de
   nieuwe naam volgens de afspraak "Wat - Van wie", bv. "Factuur - Papierhandel Vellekens".
   Bewust zonder datum, versie of lage streepjes: Jonas vroeg het heel eenvoudig te houden.
 - Bestand 1 (Document zonder titel.docx) staat als voorbeeld ingevuld in het werkdocument (rij 1). De
   leraar doet stap 4 niet klassikaal voor (Jonas, 04-10-2026): de leerling hernoemt alle vier.
 - De bedrijven komen uit de leveringen van les 05 (TV4): Bakkerij Korstjes (7.30 uur, poort 3),
   Papierhandel Vellekens (8.15 uur), Koffiebranderij De Bonenbaas (10.30 uur) en Speelgoed Tolletje
   (13.00 uur, 3 pallets, poort 3). Alle bedrijven, namen en adressen zijn fictief.
 - De zip heeft geen mapniveau: uitpakken geeft meteen vier losse bestanden.
"""
import os
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
ROMMEL = os.path.join(HERE, "rommel")
UIT = os.path.join(HERE, "WP1_Mijn-NovaDepot-map.docx")
ZIP = os.path.join(HERE, "NovaDepot-bestanden.zip")

NAVY = RGBColor(0x2A, 0x39, 0x73)
GREY = RGBColor(0x4D, 0x55, 0x73)
GRIJS = RGBColor(0x80, 0x80, 0x80)

# oude naam -> (wat, van wie, map). Volgorde = volgorde in het werkdocument; rij 1 is de demo.
ROMMELBESTANDEN = [
    ("Document zonder titel.docx", "Leveringsbon", "Bakkerij Korstjes", "Leveringen"),
    ("factuur def nieuw (2).docx", "Factuur", "Papierhandel Vellekens", "Facturen"),
    ("lijst.xlsx", "Prijslijst", "Koffiebranderij De Bonenbaas", "Prijslijsten"),
    ("IMG_20261013_131502.png", "Foto", "Speelgoed Tolletje", "Leveringen"),
]
MAPPEN = ["Leveringen", "Facturen", "Prijslijsten"]


# ---------- hulpfuncties (zoals in les 05) ----------
def lettertype(run, naam="Arial"):
    run.font.name = naam
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), naam)


def basis_document(titel):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    rpr = st.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), "Arial")
    taal = OxmlElement("w:lang")
    taal.set(qn("w:val"), "nl-BE")
    rpr.append(taal)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.1
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)
    doc.core_properties.title = titel
    doc.core_properties.author = "Toegepaste Informatica"
    return doc


def tekst(doc, s, vet=False, klein=False, cursief=False, na=6, grootte=11, kleur=None, uitlijning=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(na)
    if uitlijning is not None:
        p.alignment = uitlijning
    r = p.add_run(s)
    lettertype(r)
    r.bold = vet
    r.italic = cursief
    r.font.size = Pt(9.5 if klein else grootte)
    if klein:
        r.font.color.rgb = GREY
    if kleur is not None:
        r.font.color.rgb = kleur
    return p


def kop(doc, s, grootte=13, voor=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(voor)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(s)
    lettertype(r)
    r.bold = True
    r.font.size = Pt(grootte)
    r.font.color.rgb = NAVY
    return p


def _voeg_in_voor(ouder, el, opvolgers):
    """Voeg el in vóór het eerste bestaande element uit opvolgers (OOXML-volgorde), anders achteraan."""
    for tag in opvolgers:
        ref = ouder.find(qn(tag))
        if ref is not None:
            ref.addprevious(el)
            return el
    ouder.append(el)
    return el


def vaste_tabel(t, breedtes_cm, randkleur="8C8C8C"):
    """Vaste kolombreedtes en dunne randen, zodat Google Documenten de tabel toont zoals bedoeld."""
    tbl = t._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = _voeg_in_voor(tblpr, OxmlElement("w:tblW"),
                             ("w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                              "w:tblLayout", "w:tblCellMar", "w:tblLook"))
    tblw.set(qn("w:w"), str(int(round(sum(breedtes_cm) * 567))))
    tblw.set(qn("w:type"), "dxa")
    randen = OxmlElement("w:tblBorders")
    for kant in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + kant)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), randkleur)
        randen.append(el)
    _voeg_in_voor(tblpr, randen, ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"))
    indeling = OxmlElement("w:tblLayout")
    indeling.set(qn("w:type"), "fixed")
    _voeg_in_voor(tblpr, indeling, ("w:tblCellMar", "w:tblLook"))
    for kolom, b in zip(tbl.tblGrid.findall(qn("w:gridCol")), breedtes_cm):
        kolom.set(qn("w:w"), str(int(round(b * 567))))
    for rij in t.rows:
        for cel, b in zip(rij.cells, breedtes_cm):
            cel.width = Cm(b)
    return t


def schaduw(cel, kleur="E4E8F6"):
    tcpr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), kleur)
    tcpr.append(shd)


def cel_tekst(cel, s, vet=False, grootte=10.5, cursief=False):
    cel.text = ""
    p = cel.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(s)
    lettertype(r)
    r.bold = vet
    r.italic = cursief
    r.font.size = Pt(grootte)


def antwoordlijnen(doc, aantal=2):
    for _ in range(aantal):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("_" * 78)
        lettertype(r)
        r.font.color.rgb = GRIJS


# ---------- 1. Het werkdocument ----------
def werkdocument():
    doc = basis_document("WP1 Mijn NovaDepot-map - werkdocument")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("WERKDOCUMENT — MIJN NOVADEPOT-MAP")
    lettertype(r)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = NAVY
    tekst(doc, "Les 06 — Mijn digitale werkplek 1  ·  Toegepaste Informatica  ·  NovaDepot", klein=True, na=8)

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    vaste_tabel(t, [7.0, 4.5, 5.5])
    for i, label in enumerate(["Naam:", "Klas:", "Datum:"]):
        cel_tekst(t.rows[0].cells[i], label + " ", vet=True, grootte=11)

    kop(doc, "Zo werk je", 12, voor=10)
    for s in [
        "1.  Op de lespagina lees je wat je moet doen, stap voor stap.",
        "2.  In Google Drive maak je je NovaDepot-map. Dat is je echte werk.",
        "3.  In dit werkdocument schrijf je op wat je deed.",
        "4.  Alleen DIT document lever je in.",
    ]:
        tekst(doc, s, na=1)

    # ---- deel 1 ----
    kop(doc, "Deel 1 — Mijn mappen (stap 2)")
    tekst(doc, "Kruis aan wat je gemaakt hebt.", klein=True)
    tekst(doc, "☐  NovaDepot   (in Mijn Drive)", vet=True, na=1)
    for m in MAPPEN:
        p = tekst(doc, "☐  " + m, na=1)
        p.paragraph_format.left_indent = Cm(1.2)

    # ---- deel 2 ----
    kop(doc, "Deel 2 — Mijn bestanden (stap 4 en 5)")
    tekst(doc, "De afspraak: Wat - Van wie. Rij 1 is een voorbeeld: geef dat bestand ook die naam in je Drive.", klein=True)
    t = doc.add_table(rows=len(ROMMELBESTANDEN) + 1, cols=5)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    vaste_tabel(t, [3.8, 2.4, 3.3, 4.6, 2.9])
    for i, k in enumerate(["Oude naam", "Wat?", "Van wie?", "Nieuwe naam", "In welke map?"]):
        cel_tekst(t.rows[0].cells[i], k, vet=True)
        schaduw(t.rows[0].cells[i])
    for i, (oud, wat, wie, mp) in enumerate(ROMMELBESTANDEN, start=1):
        cel_tekst(t.rows[i].cells[0], oud, grootte=9.5)
        if i == 1:     # het voorbeeld
            for j, w in enumerate([wat, wie, f"{wat} - {wie}", mp], start=1):
                cel_tekst(t.rows[i].cells[j], w, cursief=True, grootte=9.5)
        t.rows[i].height = Cm(1.1)

    # ---- deel 3 ----
    kop(doc, "Deel 3 — Twee vragen")
    tekst(doc, "Vraag 1. Je pakte het zip-bestand uit. Waar stonden de bestanden toen? Kruis aan.", vet=True, na=2)
    tekst(doc, "☐  op deze computer, in Downloads          ☐  in mijn Google Drive", na=6)
    tekst(doc, "En waar staan ze nu, na het uploaden?", vet=True, na=2)
    tekst(doc, "☐  op deze computer, in Downloads          ☐  in mijn Google Drive", na=8)
    tekst(doc, "Vraag 2. Waarom is Factuur - Papierhandel Vellekens een betere naam dan "
               "factuur def nieuw (2)?", vet=True)
    antwoordlijnen(doc, 2)

    # ---- zelfcontrole ----
    kop(doc, "Zelfcontrole — aankruisen vóór je inlevert")
    for s in [
        "Mijn map NovaDepot staat in Mijn Drive, met drie mappen erin.",
        "Mijn vier bestanden staan in Drive, niet alleen in Downloads.",
        "Elke naam is Wat - Van wie. Achter de naam staat nog .docx, .xlsx of .png.",
        "Elk bestand zit in de juiste map.",
        "Deel 1, 2 en 3 zijn ingevuld.",
    ]:
        tekst(doc, "☐  " + s, na=1)

    # ---- extra ----
    kop(doc, "Extra — niet verplicht")
    tekst(doc, "Alleen als je al ingeleverd hebt. Klik in de opdracht op Inleveren ongedaan maken. "
               "Lever daarna opnieuw in.", klein=True)
    tekst(doc, "a) Geef je map NovaDepot een kleur: klik met 2 vingers op de map › Ordenen › Kleur van map.", vet=True, na=2)
    tekst(doc, "b) Zet een ster bij je map: klik met 2 vingers op de map › Ordenen › Toevoegen aan Met ster.", vet=True, na=2)

    doc.save(UIT)


# ---------- 2. De vier rommelbestanden ----------
def bedrijfskop(doc, wat, wie, adres):
    """Bovenaan elk document: WAT in grote letters, daaronder VAN WIE. Dat heeft de leerling nodig."""
    tekst(doc, wat.upper(), vet=True, grootte=26, kleur=NAVY, na=2)
    tekst(doc, wie, vet=True, grootte=18, na=0)
    tekst(doc, adres, klein=True, na=14)


def rommel_leveringsbon():
    doc = basis_document("Leveringsbon")
    bedrijfskop(doc, "Leveringsbon", "Bakkerij Korstjes", "Graanstraat 12, 9000 Gent  ·  fictief bedrijf")
    tekst(doc, "Aan: NovaDepot, ontvangstzone, poort 3", na=2)
    tekst(doc, "Levering: dinsdag 13 oktober 2026, 7.30 uur", na=10)
    t = doc.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    vaste_tabel(t, [8.0, 3.0, 3.0])
    for i, k in enumerate(["Product", "Aantal", "Pallets"]):
        cel_tekst(t.rows[0].cells[i], k, vet=True)
        schaduw(t.rows[0].cells[i])
    for i, rij in enumerate([("Volkorenbrood, gesneden", "240", "1"), ("Pistolets, diepvries", "600", "1"),
                             ("Totaal", "", "2")], start=1):
        for j, w in enumerate(rij):
            cel_tekst(t.rows[i].cells[j], w, vet=(i == 3))
    tekst(doc, "")
    tekst(doc, "Ontvangen door: ............................................", na=2)
    doc.save(os.path.join(ROMMEL, ROMMELBESTANDEN[0][0]))


def rommel_factuur():
    doc = basis_document("Factuur")
    bedrijfskop(doc, "Factuur", "Papierhandel Vellekens", "Bladzijdelaan 7, 9050 Gentbrugge  ·  fictief bedrijf")
    tekst(doc, "Factuurnummer: 2026-0417", na=2)
    tekst(doc, "Datum: 13 oktober 2026", na=2)
    tekst(doc, "Klant: NovaDepot, Magazijnweg 1, Gent", na=10)
    t = doc.add_table(rows=4, cols=4)
    t.style = "Table Grid"
    vaste_tabel(t, [7.0, 2.2, 2.6, 2.6])
    for i, k in enumerate(["Product", "Aantal", "Prijs", "Totaal"]):
        cel_tekst(t.rows[0].cells[i], k, vet=True)
        schaduw(t.rows[0].cells[i])
    for i, rij in enumerate([("Kopieerpapier A4, doos van 5 pakken", "40", "€ 24,00", "€ 960,00"),
                             ("Etiketten voor pallets, rol", "10", "€ 8,50", "€ 85,00"),
                             ("Totaal zonder btw", "", "", "€ 1 045,00")], start=1):
        for j, w in enumerate(rij):
            cel_tekst(t.rows[i].cells[j], w, vet=(i == 3))
    tekst(doc, "")
    tekst(doc, "Te betalen binnen 30 dagen.", klein=True)
    doc.save(os.path.join(ROMMEL, ROMMELBESTANDEN[1][0]))


def rommel_prijslijst():
    wb = Workbook()
    ws = wb.active
    ws.title = "Prijslijst"
    ws["A1"] = "PRIJSLIJST"
    ws["A1"].font = Font(name="Arial", size=22, bold=True, color="2A3973")
    ws["A2"] = "Koffiebranderij De Bonenbaas"
    ws["A2"].font = Font(name="Arial", size=16, bold=True)
    ws["A3"] = "fictief bedrijf  ·  prijzen voor groothandels, oktober 2026"
    ws["A3"].font = Font(name="Arial", size=9, color="4D5573")
    rand = Border(*[Side(style="thin", color="8C8C8C")] * 4)
    kop = ["Product", "Inhoud", "Prijs per stuk"]
    for j, k in enumerate(kop, start=1):
        c = ws.cell(5, j, k)
        c.font = Font(name="Arial", size=11, bold=True)
        c.fill = PatternFill("solid", fgColor="E4E8F6")
        c.border = rand
    rijen = [("Koffiebonen Classic", "1 kg", 12.50), ("Koffiebonen Espresso", "1 kg", 14.20),
             ("Gemalen koffie", "500 g", 6.80), ("Koffiepads", "36 stuks", 4.95)]
    for i, (prod, inh, prijs) in enumerate(rijen, start=6):
        for j, w in enumerate((prod, inh, prijs), start=1):
            c = ws.cell(i, j, w)
            c.font = Font(name="Arial", size=11)
            c.border = rand
            if j == 3:
                c.number_format = '€ #,##0.00'
                c.alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 16
    wb.save(os.path.join(ROMMEL, ROMMELBESTANDEN[2][0]))


def rommel_foto():
    """De naam die de camera zelf gaf. Op de foto: een pallet met het etiket van Speelgoed Tolletje."""
    from PIL import Image, ImageDraw, ImageFont

    def lettertype_pil(grootte):
        for pad in ("/System/Library/Fonts/Supplemental/Arial Bold.ttf",
                    "/System/Library/Fonts/Supplemental/Arial.ttf",
                    "/Library/Fonts/Arial.ttf",
                    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"):
            if os.path.exists(pad):
                try:
                    return ImageFont.truetype(pad, grootte)
                except Exception:
                    pass
        return ImageFont.load_default(size=grootte)

    B, H = 1200, 800
    img = Image.new("RGB", (B, H), (214, 216, 210))          # grijze magazijnvloer
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, B, 260], fill=(236, 238, 234))         # muur
    d.rectangle([60, 40, 420, 110], fill=(42, 57, 115))       # bord van de poort
    d.text((85, 55), "POORT 3", font=lettertype_pil(44), fill=(255, 255, 255))

    # de pallet met dozen
    hout = (176, 132, 82)
    d.rectangle([300, 640, 900, 690], fill=hout)
    for x in range(310, 900, 120):
        d.rectangle([x, 690, x + 60, 720], fill=(150, 110, 66))
    karton = [(205, 160, 105), (195, 150, 98), (210, 168, 112)]
    for rij in range(3):
        for kol in range(4):
            x0, y0 = 310 + kol * 148, 460 - rij * 150 + 30
            d.rectangle([x0, y0, x0 + 140, y0 + 145], fill=karton[(rij + kol) % 3], outline=(150, 115, 70), width=3)

    # het etiket: WIE staat er groot op
    d.rectangle([420, 300, 900, 520], fill=(255, 255, 255), outline=(60, 60, 60), width=4)
    d.text((445, 315), "Speelgoed Tolletje", font=lettertype_pil(46), fill=(160, 55, 42))
    d.text((445, 385), "3 pallets  -  poort 3", font=lettertype_pil(34), fill=(40, 40, 40))
    d.text((445, 440), "13-10-2026  13.00 uur", font=lettertype_pil(30), fill=(77, 85, 115))
    d.text((60, 740), "foto bij ontvangst  -  fictief bedrijf", font=lettertype_pil(26), fill=(110, 110, 105))
    img.save(os.path.join(ROMMEL, ROMMELBESTANDEN[3][0]), "PNG")


# ---------- 3. De zip voor Google Classroom ----------
def maak_zip():
    # zonder mapniveau: uitpakken geeft meteen vier losse bestanden
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for naam, *_ in ROMMELBESTANDEN:
            z.write(os.path.join(ROMMEL, naam), arcname=naam)


if __name__ == "__main__":
    os.makedirs(ROMMEL, exist_ok=True)
    werkdocument()
    rommel_leveringsbon()
    rommel_factuur()
    rommel_prijslijst()
    rommel_foto()
    maak_zip()
    print("Klaar.")
    print("  werkdocument om in te leveren : " + os.path.basename(UIT))
    print("  aan de opdracht te hangen     : " + os.path.basename(ZIP))
    print("  losse bronbestanden           : rommel/ ({} bestanden)".format(len(ROMMELBESTANDEN)))
